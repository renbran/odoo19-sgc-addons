# -*- coding: utf-8 -*-
"""Upload buckets and the accepted column formats for the client portal.

The portal intentionally uses a small, fixed vocabulary.  A bucket is both the
portal category and the ``doc_type`` stored on ``sgc.accounting.upload`` so old
uploads keep working.  Keep the CSV sample and the required list in sync: the
sample download is generated from the same definition that validates uploads.
"""

import csv
import io
import os
import re

from odoo import _

BUCKETS = (
    {
        "code": "customer_invoices",
        "name": "Sales Invoices",
        "short": "Invoices you issue to customers",
        "description": (
            "Sales invoices, credit notes issued to customers and invoice "
            "supporting schedules."
        ),
        "formats": "CSV, XLSX, XLS, PDF",
        "required": (
            "Invoice Number",
            "Invoice Date",
            "Customer Name",
            "Currency",
            "Description",
            "Quantity",
            "Unit Price",
            "Tax Rate",
            "Tax Amount",
            "Total",
            "Due Date",
        ),
        "example": (
            "INV-2026-0001",
            "2026-10-01",
            "Example Customer LLC",
            "AED",
            "Consulting services",
            "1",
            "1000.00",
            "5",
            "50.00",
            "1050.00",
            "2026-10-31",
        ),
        "notes": (
            "Use one row per invoice line; repeat Invoice Number for multiple lines.",
            "Dates must be YYYY-MM-DD; amounts use a dot decimal separator.",
        ),
    },
    {
        "code": "vendor_bills",
        "name": "Purchase Bills",
        "short": "Bills and expenses from suppliers",
        "description": (
            "Supplier invoices, utility bills, rent invoices and other "
            "purchase documents."
        ),
        "formats": "CSV, XLSX, XLS, PDF",
        "required": (
            "Bill Number",
            "Bill Date",
            "Vendor Name",
            "Currency",
            "Description",
            "Quantity",
            "Unit Price",
            "Tax Rate",
            "Tax Amount",
            "Total",
            "Due Date",
        ),
        "example": (
            "BILL-2026-0001",
            "2026-10-02",
            "Example Supplier LLC",
            "AED",
            "Office supplies",
            "1",
            "500.00",
            "5",
            "25.00",
            "525.00",
            "2026-10-31",
        ),
        "notes": (
            "Use one row per bill line; repeat Bill Number for multiple lines.",
            "The vendor TRN is optional but improves matching when present.",
        ),
    },
    {
        "code": "payments",
        "name": "Payments",
        "short": "Money received or paid",
        "description": (
            "Customer receipts, supplier payments, refunds and other payment "
            "advice."
        ),
        "formats": "CSV, XLSX, XLS, PDF",
        "required": (
            "Payment Date",
            "Direction",
            "Reference",
            "Party",
            "Invoice Number",
            "Amount",
            "Currency",
            "Bank Account",
            "Payment Method",
        ),
        "example": (
            "2026-10-05",
            "Inbound",
            "PAY-2026-0001",
            "Example Customer LLC",
            "INV-2026-0001",
            "1050.00",
            "AED",
            "Main Operating Account",
            "Bank transfer",
        ),
        "notes": (
            "Direction must be Inbound or Outbound.",
            "Invoice Number may be blank for an on-account payment.",
        ),
    },
    {
        "code": "bank_statements",
        "name": "Bank Statements",
        "short": "Bank or card statement lines",
        "description": (
            "Bank statements, card statements and other cash movement extracts."
        ),
        "formats": "CSV, XLSX, XLS, PDF",
        "required": (
            "Statement Date",
            "Description",
            "Reference",
            "Debit",
            "Credit",
            "Balance",
            "Currency",
            "Bank Account",
        ),
        "example": (
            "2026-10-05",
            "Customer receipt",
            "PAY-2026-0001",
            "",
            "1050.00",
            "1050.00",
            "AED",
            "Main Operating Account",
        ),
        "notes": (
            "Use either Debit or Credit on each line; leave the other blank or zero.",
            "Do not include header/footer rows inside the data area.",
        ),
    },
    {
        "code": "payroll_journals",
        "name": "Payroll",
        "short": "Payroll and employee cost journals",
        "description": (
            "Monthly payroll journals, salary summaries and employee cost "
            "schedules."
        ),
        "formats": "CSV, XLSX, XLS, PDF",
        "required": (
            "Payroll Date",
            "Employee Code",
            "Employee Name",
            "Basic Salary",
            "Allowances",
            "Deductions",
            "Net Pay",
            "Currency",
        ),
        "example": (
            "2026-10-31",
            "EMP-001",
            "Example Employee",
            "8000.00",
            "2000.00",
            "500.00",
            "9500.00",
            "AED",
        ),
        "notes": (
            "Do not send bank account numbers or Emirates ID numbers in this bucket.",
            "One row per employee per payroll period.",
        ),
    },
    {
        "code": "credit_notes",
        "name": "Credit Notes",
        "short": "Sales or purchase credit notes",
        "description": (
            "Customer refunds, supplier credit notes and debit-note adjustments."
        ),
        "formats": "CSV, XLSX, XLS, PDF",
        "required": (
            "Credit Note Number",
            "Credit Note Date",
            "Party Type",
            "Party Name",
            "Original Invoice",
            "Currency",
            "Description",
            "Quantity",
            "Unit Price",
            "Tax Rate",
            "Tax Amount",
            "Total",
        ),
        "example": (
            "CN-2026-0001",
            "2026-10-10",
            "Customer",
            "Example Customer LLC",
            "INV-2026-0001",
            "AED",
            "Service adjustment",
            "1",
            "100.00",
            "5",
            "5.00",
            "105.00",
        ),
        "notes": (
            "Party Type must be Customer or Supplier.",
            "Reference the original invoice where possible.",
        ),
    },
    {
        "code": "expense_receipts",
        "name": "Expense Receipts",
        "short": "Petty cash and reimbursement receipts",
        "description": (
            "Employee expense claims, petty cash receipts and small supplier "
            "receipts."
        ),
        "formats": "CSV, XLSX, XLS, PDF",
        "required": (
            "Receipt Date",
            "Vendor Name",
            "Category",
            "Description",
            "Amount",
            "Tax Amount",
            "Total",
            "Currency",
            "Payment Method",
            "Receipt Number",
        ),
        "example": (
            "2026-10-12",
            "Example Vendor",
            "Travel",
            "Taxi to client meeting",
            "100.00",
            "5.00",
            "105.00",
            "AED",
            "Cash",
            "RCPT-2026-0001",
        ),
        "notes": (
            "Attach the receipt image or PDF separately when it is available.",
            "Category is free text but should be a consistent expense category.",
        ),
    },
    {
        "code": "master_data",
        "name": "Master Data",
        "short": "Customers, suppliers and opening balances",
        "description": (
            "Customer and supplier lists, chart-of-accounts mappings and "
            "opening balances."
        ),
        "formats": "CSV, XLSX, XLS",
        "required": (
            "Record Type",
            "Code",
            "Name",
            "TRN",
            "Email",
            "Phone",
            "Address",
            "Currency",
        ),
        "example": (
            "Customer",
            "CUST-0001",
            "Example Customer LLC",
            "100000000000003",
            "ap@example.com",
            "+971500000000",
            "Dubai, UAE",
            "AED",
        ),
        "notes": (
            "Record Type must be Customer, Supplier, Account or Opening Balance.",
            "TRN must be 15 digits when supplied.",
        ),
    },
    {
        "code": "other",
        "name": "Other Documents",
        "short": "Anything that does not fit another bucket",
        "description": (
            "Contracts, statements, correspondence and other supporting "
            "documents."
        ),
        "formats": "CSV, XLSX, XLS, PDF, JPG, JPEG, PNG",
        "required": (
            "Document Type",
            "Reference",
            "Document Date",
            "Party Name",
            "Amount",
            "Currency",
            "Description",
        ),
        "example": (
            "Contract",
            "CTR-2026-0001",
            "2026-10-15",
            "Example Customer LLC",
            "",
            "AED",
            "Signed service agreement",
        ),
        "notes": (
            "Use this bucket only when no accounting bucket fits.",
            "A reviewer will classify the document before processing.",
        ),
    },
)

BUCKETS_BY_CODE = {bucket["code"]: bucket for bucket in BUCKETS}
DOC_TYPES = [(bucket["code"], bucket["name"]) for bucket in BUCKETS]


def bucket(code):
    return BUCKETS_BY_CODE.get(code)


def bucket_name(code):
    item = bucket(code)
    return item["name"] if item else _("Document")


def _normalise_header(value):
    text = (value or "").strip().lower()
    text = text.replace("_", " ").replace("-", " ")
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


# Common customer spellings are accepted, but the canonical name is always
# shown in the downloadable sample.  This avoids rejecting a valid file simply
# because a supplier used "Invoice No." instead of "Invoice Number".
HEADER_ALIASES = {
    "invoice no": "invoice number",
    "invoice #": "invoice number",
    "invoice": "invoice number",
    "bill no": "bill number",
    "bill #": "bill number",
    "bill": "bill number",
    "credit note no": "credit note number",
    "credit note #": "credit note number",
    "customer": "customer name",
    "client": "customer name",
    "client name": "customer name",
    "vendor": "vendor name",
    "supplier": "vendor name",
    "supplier name": "vendor name",
    "party": "party name",
    "total amount": "total",
    "gross total": "total",
    "vat amount": "tax amount",
    "vat": "tax amount",
    "tax": "tax amount",
    "currency code": "currency",
}

DATE_ALIASES = {
    "customer_invoices": "invoice date",
    "vendor_bills": "bill date",
    "payments": "payment date",
    "bank_statements": "statement date",
    "payroll_journals": "payroll date",
    "credit_notes": "credit note date",
    "expense_receipts": "receipt date",
    "other": "document date",
}

FORMAT_EXTENSIONS = {
    "csv": ".csv",
    "xlsx": ".xlsx",
    "xls": ".xls",
    "pdf": ".pdf",
    "jpg": ".jpg",
    "jpeg": ".jpeg",
    "png": ".png",
}

# The intake agent reads and books at most this many data rows per upload;
# larger files must be split. Keep in sync with the portal texts that state
# the limit so customers are never surprised by a rejection.
MAX_UPLOAD_ROWS = 50


def canonical_headers(values, bucket_code=None):
    headers = []
    for value in values:
        normalised = _normalise_header(value)
        if normalised == "date" and bucket_code in DATE_ALIASES:
            headers.append(DATE_ALIASES[bucket_code])
        else:
            headers.append(HEADER_ALIASES.get(normalised, normalised))
    return headers


def _check_headers(bucket_code, headers):
    definition = bucket(bucket_code)
    if not definition:
        return ["Unknown bucket"]
    present = set(headers)
    # Compare through the same alias chain used for uploaded files, so a
    # required label such as "Party" (payments) canonicalises exactly like the
    # header row of the downloadable sample does.
    return [
        label
        for label in definition["required"]
        if canonical_headers([label], bucket_code)[0] not in present
    ]


def _csv_headers(bucket_code, data):
    text = data.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = [row for row in reader if any((cell or "").strip() for cell in row)]
    if not rows:
        return [], 0
    return canonical_headers(rows[0], bucket_code), max(len(rows) - 1, 0)


def _xlsx_headers(bucket_code, data):
    try:
        import openpyxl
    except ImportError:
        return None, 0, "XLSX validation is unavailable on this server."
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        first = next(rows, None)
        if not first:
            workbook.close()
            return [], 0, None
        headers = canonical_headers(first, bucket_code)
        count = sum(
            1
            for row in rows
            if any(str(cell).strip() for cell in row if cell is not None)
        )
        workbook.close()
        return headers, count, None
    except Exception:
        return None, 0, "The XLSX file could not be read."


def _xls_headers(bucket_code, data):
    try:
        import xlrd
    except ImportError:
        return None, 0, "XLS validation is unavailable on this server."
    try:
        workbook = xlrd.open_workbook(file_contents=data)
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows == 0:
            return [], 0, None
        headers = canonical_headers(sheet.row_values(0), bucket_code)
        return headers, max(sheet.nrows - 1, 0), None
    except Exception:
        return None, 0, "The XLS file could not be read."


def validate_upload_format(bucket_code, filename, data):
    """Return a format-validation result.

    ``error`` is a user-facing string or ``None``.  PDFs and images are accepted
    without structural validation because their contents are extracted by the
    agent.  CSV/XLSX/XLS files are checked against the bucket's required
    columns before the upload is accepted.
    """
    definition = bucket(bucket_code)
    if not definition:
        return {"error": _("Select a valid upload bucket."), "missing": []}
    extension = os.path.splitext(filename or "")[1].lower()
    allowed_formats = {
        item.strip().lower() for item in definition["formats"].split(",")
    }
    allowed_extensions = {
        FORMAT_EXTENSIONS[item]
        for item in allowed_formats
        if item in FORMAT_EXTENSIONS
    }
    if extension not in allowed_extensions:
        return {
            "error": _(
                "%(bucket)s accepts only %(formats)s files."
            )
            % {
                "bucket": definition["name"],
                "formats": definition["formats"],
            },
            "missing": [],
        }
    if extension in {".pdf", ".jpg", ".jpeg", ".png"}:
        return {"error": None, "missing": [], "row_count": None}

    if extension == ".csv":
        try:
            headers, row_count = _csv_headers(bucket_code, data)
            parse_error = None
        except UnicodeDecodeError:
            return {
                "error": _("CSV files must be UTF-8 encoded."),
                "missing": [],
            }
        except csv.Error:
            return {"error": _("The CSV file could not be read."), "missing": []}
    elif extension == ".xlsx":
        headers, row_count, parse_error = _xlsx_headers(bucket_code, data)
    else:
        headers, row_count, parse_error = _xls_headers(bucket_code, data)

    if parse_error:
        return {"error": parse_error, "missing": []}
    if not headers:
        return {"error": _("The file has no header row."), "missing": []}
    missing = _check_headers(bucket_code, headers)
    if missing:
        return {
            "error": _("The file is missing required columns."),
            "missing": missing,
            "headers": headers,
        }
    if not row_count:
        return {"error": _("Add at least one data row below the header."), "missing": []}
    if row_count > MAX_UPLOAD_ROWS:
        return {
            "error": _(
                "The file has %(rows)s data rows but the maximum per upload is "
                "%(limit)s. Please split the file into smaller batches of up to "
                "%(limit)s rows and upload each one."
            )
            % {"rows": row_count, "limit": MAX_UPLOAD_ROWS},
            "missing": [],
        }
    return {
        "error": None,
        "missing": [],
        "row_count": row_count,
        "headers": headers,
    }


def sample_csv(bucket_code):
    definition = bucket(bucket_code)
    if not definition:
        return None
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(definition["required"])
    writer.writerow(definition["example"])
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def find_column(headers, label):
    """Return the header that matches ``label``, or ``None``.

    Matching is case-insensitive and ignores punctuation, so ``Total`` finds
    ``Total Amount`` only when the labels normalise to the same value.
    """
    target = _normalise_header(label)
    for header in headers or ():
        if _normalise_header(header) == target:
            return header
    return None


def _cell_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_from_pairs(headers, rows, limit):
    collected = []
    for row in rows[:limit]:
        record = {}
        for index, header in enumerate(headers):
            value = row[index] if index < len(row) else ""
            record[header] = _cell_value(value)
        collected.append(record)
    return {
        "headers": list(headers),
        "rows": collected,
        "truncated": len(rows) > limit,
    }


def read_rows(bucket_code, filename, data, limit=50):
    """Return the parsed source rows of a tabular upload.

    PDF and image uploads have no tabular structure and return ``None``.
    CSV, XLSX and XLS uploads return a dict with ``headers``, ``rows`` and
    ``truncated`` keys. The rows are the authoritative source values used by
    the deterministic checksum gate, so they are never taken from an agent
    reply.
    """
    extension = os.path.splitext(filename or "")[1].lower()
    if extension == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except (UnicodeDecodeError, AttributeError):
            return None
        try:
            reader = csv.reader(io.StringIO(text))
            rows = [row for row in reader if any((cell or "").strip() for cell in row)]
        except csv.Error:
            return None
        if not rows:
            return None
        return _rows_from_pairs(canonical_headers(rows[0], bucket_code), rows[1:], limit)
    if extension == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            return None
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True
            )
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            first = next(values, None)
            if not first:
                workbook.close()
                return None
            headers = canonical_headers(first, bucket_code)
            collected = [
                list(row)
                for row in values
                if any(str(cell).strip() for cell in row if cell is not None)
            ]
            workbook.close()
            return _rows_from_pairs(headers, collected, limit)
        except Exception:
            return None
    if extension == ".xls":
        try:
            import xlrd
        except ImportError:
            return None
        try:
            workbook = xlrd.open_workbook(file_contents=data)
            sheet = workbook.sheet_by_index(0)
            if sheet.nrows == 0:
                return None
            headers = canonical_headers(sheet.row_values(0), bucket_code)
            collected = [sheet.row_values(index) for index in range(1, sheet.nrows)]
            return _rows_from_pairs(headers, collected, limit)
        except Exception:
            return None
    return None
